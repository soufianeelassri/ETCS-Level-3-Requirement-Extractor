"""
Step 6: Write the Excel file. 
Figures and tables go inside the Description cell, under the text.
"""

import io
import math
from pathlib import Path
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font
from openpyxl.utils.units import pixels_to_EMU

from describer import Requirement, render_media

COLUMNS = [("Topic", 30), ("Requirement ID", 24), ("Description", 85), ("Traceability", 36)]
DESCRIPTION_COL = 3
ROW_MAX_PT = 409  # Excel won't go higher
LINE_PT = 15  # Calibri 11

def text_height(text: str, width_chars: float) -> float:
    # rough estimate, Excel doesn't resize rows written by openpyxl
    lines = sum(max(1, math.ceil(len(l) * 1.1 / width_chars)) for l in text.split("\n"))
    return lines * LINE_PT + 4

def write_requirements(requirements: list[Requirement], pdf_path: Path, output: Path) -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Requirements"
    for col, (name, width) in enumerate(COLUMNS, start=1):
        ws.cell(1, col, name).font = Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    desc_width = COLUMNS[DESCRIPTION_COL - 1][1]
    desc_px = int(desc_width * 7 + 5) - 12
    images = 0
    for row, r in enumerate(requirements, start=2):
        values = [r.topic, r.id, r.description, r.traceability]
        for col, value in enumerate(values, start=1):
            ws.cell(row, col, value).alignment = Alignment(wrap_text=True, vertical="top")
        height = max(text_height(str(v), w) for v, (_, w) in zip(values, COLUMNS))

        # images start under the text and shrink to fit the row height limit
        y = text_height(r.description, desc_width) if r.media else height
        room = ROW_MAX_PT - y - 4 * len(r.media)
        for m in r.media:
            img = XLImage(io.BytesIO(render_media(pdf_path, m)))
            scale = min(desc_px / img.width, room / len(r.media) / (img.height * 0.75), 1.0)
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            marker = AnchorMarker(col=DESCRIPTION_COL - 1, colOff=pixels_to_EMU(6),
                                  row=row - 1, rowOff=pixels_to_EMU(int(y / 0.75)))
            img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(
                pixels_to_EMU(img.width), pixels_to_EMU(img.height)))
            ws.add_image(img)
            y += img.height * 0.75 + 4
            images += 1
        ws.row_dimensions[row].height = min(max(height, y + 4), ROW_MAX_PT)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return images